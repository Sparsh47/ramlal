import os
import sys
import openpyxl
from datetime import datetime

# Add the root directory to the python path so we can import from lib
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

from lib.db import SessionLocal, Job, get_existing_urls_db

def migrate_excel_to_db(filename: str = "jobs.xlsx"):
    if not os.path.exists(filename):
        print(f"File {filename} not found.")
        return

    print(f"Loading data from {filename}...")
    wb = openpyxl.load_workbook(filename)
    if "Job Listings" not in wb.sheetnames:
        print("Sheet 'Job Listings' not found in excel file.")
        return
        
    ws = wb["Job Listings"]
    
    # Map headers to column indices
    headers = {cell.value: cell.column for cell in ws[1]}
    
    required_headers = ["Score", "Title", "Company", "URL"]
    if not all(h in headers for h in required_headers):
        print(f"Missing required headers. Found: {list(headers.keys())}")
        return

    session = SessionLocal()
    existing_urls = get_existing_urls_db()
    added_count = 0
    
    print("Migrating jobs to database...")
    try:
        for row in range(2, ws.max_row + 1):
            url = ws.cell(row=row, column=headers["URL"]).value
            
            if not url or url in existing_urls:
                continue
                
            score_val = ws.cell(row=row, column=headers["Score"]).value
            score = float(score_val) if score_val is not None else 0.0
            
            auto_ready_val = ws.cell(row=row, column=headers.get("Auto Apply Ready", -1)).value
            auto_ready = True if auto_ready_val and "✅" in str(auto_ready_val) else False
            
            scored_on = ws.cell(row=row, column=headers.get("Scored On", -1)).value or "unknown"
            title = ws.cell(row=row, column=headers["Title"]).value or ""
            company = ws.cell(row=row, column=headers["Company"]).value or ""
            reasons = ws.cell(row=row, column=headers.get("Reasons", -1)).value or ""
            
            new_job = Job(
                score=score,
                auto_apply_ready=auto_ready,
                scored_on=scored_on,
                title=title,
                company=company,
                reasons=reasons,
                url=url,
                date_found=datetime.utcnow()
            )
            session.add(new_job)
            existing_urls.add(url)
            added_count += 1
            
        session.commit()
        print(f"Successfully migrated {added_count} new jobs to Postgres database.")
    except Exception as e:
        session.rollback()
        print(f"Database error: {e}")
    finally:
        session.close()

if __name__ == "__main__":
    migrate_excel_to_db()
