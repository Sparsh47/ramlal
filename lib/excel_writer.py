import os
import openpyxl
from datetime import date

def get_existing_urls(filename: str = "jobs.xlsx") -> set[str]:
    """Returns a set of URLs already present in the Excel file."""
    if not os.path.exists(filename):
        return set()
    
    wb = openpyxl.load_workbook(filename)
    if "Job Listings" not in wb.sheetnames:
        return set()
    
    ws = wb["Job Listings"]
    urls = set()
    
    # Assuming URL is the 7th column (index 6, but openpyxl is 1-indexed so col 7)
    # Let's dynamically find the URL column index from headers just in case
    url_col_idx = None
    for cell in ws[1]:
        if cell.value == "URL":
            url_col_idx = cell.column
            break
            
    if url_col_idx is None:
        return set()
        
    for row in range(2, ws.max_row + 1):
        cell_val = ws.cell(row=row, column=url_col_idx).value
        if cell_val:
            urls.add(str(cell_val).strip())
            
    return urls


def save_to_excel(jobs: list[dict], filename: str = "jobs.xlsx"):
    file_exists = os.path.exists(filename)
    
    if file_exists:
        wb = openpyxl.load_workbook(filename)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Job Listings"
        headers = ["Score", "Auto Apply Ready", "Scored On", "Title", "Company", "Reasons", "URL", "Date Found"]
        ws.append(headers)
        bold_font = openpyxl.styles.Font(bold=True)
        for cell in ws[1]:
            cell.font = bold_font

    # Styles
    green_fill = openpyxl.styles.PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    yellow_fill = openpyxl.styles.PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

    existing_urls = get_existing_urls(filename)
    added_count = 0

    for job in jobs:
        url = job.get("url")
        if url in existing_urls:
            continue
            
        auto_ready = job.get("auto_apply_ready", False)
        row_data = [
            job.get("score"),
            "✅ Yes" if auto_ready else "❌ Manual",
            job.get("scored_on", "unknown"),
            job.get("title"),
            job.get("company"),
            job.get("reasons"),
            url,
            str(date.today())
        ]
        ws.append(row_data)
        added_count += 1
        existing_urls.add(url)
        
        # Color code the row
        current_row = ws.max_row
        fill = green_fill if auto_ready else yellow_fill
        for cell in ws[current_row]:
            cell.fill = fill

    # Auto-width columns
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 60)

    wb.save(filename)
    print(f"Appended {added_count} new jobs to {filename} (Skipped {len(jobs) - added_count} duplicates)")